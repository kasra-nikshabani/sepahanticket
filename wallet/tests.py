"""تست‌های جریان پول در برداشت وجه.

چرا این فایل وجود دارد: هر مسیری که پول را جابه‌جا می‌کند باید ثابت کند که
پول نه گم می‌شود و نه دو بار شمرده می‌شود. اشتباه در این مسیر مثل اشتباه در
یک صفحه‌ی نمایشی نیست -- مستقیم روی حساب بانکی آدم‌ها اثر می‌گذارد.
"""
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase

from .models import (Wallet, Transaction, WithdrawalRequest,
                     get_withdrawable_amount, is_valid_iban, normalize_iban,
                     WITHDRAWABLE_PREFIXES)

User = get_user_model()

VALID_IBAN = 'IR270170000000100324200001'


class WithdrawableAmountTests(TestCase):
    """«چقدر قابل برداشت است» -- قاعده‌ای که کل ویژگی رویش سوار است."""

    def setUp(self):
        self.user = User.objects.create_user(username='u1', password='x')
        # کیف پول با سیگنال post_save روی User ساخته می‌شود، پس فقط برش می‌داریم.
        self.wallet = Wallet.objects.get(user=self.user)

    def _credit(self, amount, reference_id, tx_type='deposit'):
        self.wallet.add_balance(amount=amount, reference_id=reference_id, tx_type=tx_type)

    def test_self_charged_money_is_not_withdrawable(self):
        """پولی که کاربر خودش از درگاه شارژ کرده قابل برداشت نیست."""
        self._credit(5_000_000, reference_id='4734693545')  # track_id درگاه
        self.assertEqual(self.wallet.balance, 5_000_000)
        self.assertEqual(get_withdrawable_amount(self.user), 0)

    def test_compensation_money_is_withdrawable(self):
        self._credit(3_000_000, reference_id='compensate-64-12')
        self.assertEqual(get_withdrawable_amount(self.user), 3_000_000)

    def test_all_compensation_prefixes_count(self):
        for i, pre in enumerate(WITHDRAWABLE_PREFIXES):
            self._credit(1_000_000, reference_id=f'{pre}{i}')
        self.assertEqual(get_withdrawable_amount(self.user),
                         1_000_000 * len(WITHDRAWABLE_PREFIXES))

    def test_spending_comes_out_of_the_compensation_share_first(self):
        """خرجِ بلیط اول از سهمِ جبرانی کم می‌شود، نه از پولِ شارژیِ کاربر.

        نسخه‌ی اول این تست عکسِ این را ثبت کرده بود («خرج اول از پول شارژی
        کم می‌شود»)، که به نظر به نفع کاربر می‌آمد ولی راهی باز می‌گذاشت تا
        پولِ شارژی از کیف پول بیرون کشیده شود. آن قاعده -- و این تست -- عوض
        شدند؛ توضیح کامل در get_withdrawable_amount.
        """
        self._credit(5_000_000, reference_id='compensate-64-1')
        self._credit(1_000_000, reference_id='4734693545')      # شارژ شخصی
        self.wallet.deduct_balance(amount=2_000_000, reference_id='order-1',
                                   tx_type='ticket_purchase')
        self.assertEqual(self.wallet.balance, 4_000_000)
        # ۵ جبرانی منهای ۲ خرج = ۳؛ آن یک میلیونِ شارژی دست‌نخورده قفل می‌ماند.
        self.assertEqual(get_withdrawable_amount(self.user), 3_000_000)

    def test_spending_the_compensation_then_charging_does_not_reopen_withdrawal(self):
        """راهِ دورِ بیرون‌کشیدنِ پولِ شارژی -- باید بسته باشد.

        سناریو: کاربر ۳ میلیون جبرانی می‌گیرد، با همان بلیط می‌خرد، بعد
        ۳ میلیون از کارت خودش شارژ می‌کند. اگر خرج از سهمِ جبرانی کم نشود،
        همان پولِ شارژی «قابل برداشت» شمرده می‌شود و سایت تبدیل به کانال
        انتقال پول می‌شود: ورود با کارتِ الف، خروج به حسابِ ب.
        """
        self._credit(3_000_000, reference_id='compensate-64-7')
        self.wallet.deduct_balance(amount=3_000_000, reference_id='order-9',
                                   tx_type='ticket_purchase')
        self._credit(3_000_000, reference_id='4734699999')      # شارژ از درگاه
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 3_000_000)
        self.assertEqual(get_withdrawable_amount(self.user), 0,
                         'پولِ شارژیِ خود کاربر قابل برداشت شد')

    def test_partly_spent_compensation_is_withdrawable_only_for_the_rest(self):
        self._credit(3_000_000, reference_id='compensate-64-8')
        self.wallet.deduct_balance(amount=1_000_000, reference_id='order-10',
                                   tx_type='ticket_purchase')
        self.assertEqual(get_withdrawable_amount(self.user), 2_000_000)

    def test_own_charge_stays_locked_even_next_to_compensation(self):
        """۳ جبرانی + ۵ شارژِ خودش، ۵ بلیط می‌خرد -> باقی‌مانده مالِ خودش است."""
        self._credit(3_000_000, reference_id='compensate-64-11')
        self._credit(5_000_000, reference_id='4734698888')
        self.wallet.deduct_balance(amount=5_000_000, reference_id='order-11',
                                   tx_type='ticket_purchase')
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 3_000_000)
        self.assertEqual(get_withdrawable_amount(self.user), 0)

    def test_bulk_total_matches_the_per_user_rule(self):
        """عدد صفحه‌ی تنظیمات باید دقیقاً با جمعِ تک‌تک کاربران یکی باشد."""
        from .models import withdrawable_totals
        other = User.objects.create_user(username='u1b', password='x')
        Wallet.objects.get(user=other).add_balance(
            amount=2_000_000, reference_id='compensate-64-99')
        self._credit(3_000_000, reference_id='compensate-64-98')
        self._credit(1_000_000, reference_id='4734697777')       # شارژ شخصی
        users, total = withdrawable_totals()
        expected = get_withdrawable_amount(self.user) + get_withdrawable_amount(other)
        self.assertEqual(total, expected)
        self.assertEqual(total, 5_000_000)                        # شارژ شخصی نیامده
        self.assertEqual(users, 2)

    def test_never_exceeds_actual_balance(self):
        """اگر پول جبرانی خرج بلیط شده باشد، دیگر قابل برداشت نیست."""
        self._credit(5_000_000, reference_id='compensate-64-1')
        self.wallet.deduct_balance(amount=4_500_000, reference_id='order-1',
                                   tx_type='ticket_purchase')
        self.assertEqual(get_withdrawable_amount(self.user), 500_000)


class WithdrawalFlowTests(TestCase):
    """چرخه‌ی کامل: ثبت -> تأیید -> واریز، و مسیر رد شدن."""

    def setUp(self):
        self.user = User.objects.create_user(username='u2', password='x')
        self.admin = User.objects.create_user(username='adm', password='x', is_staff=True)
        # کیف پول با سیگنال post_save روی User ساخته می‌شود، پس فقط برش می‌داریم.
        self.wallet = Wallet.objects.get(user=self.user)
        self.wallet.add_balance(amount=5_000_000, reference_id='compensate-64-9',
                                tx_type='deposit')

    def test_request_holds_the_money_immediately(self):
        """پول باید همان لحظه کسر شود، نه موقع واریز."""
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 2_000_000)
        self.assertEqual(req.status, 'pending')
        tx = Transaction.objects.get(reference_id=f'WD-{req.pk}')
        self.assertEqual(tx.amount, -3_000_000)
        self.assertEqual(tx.transaction_type, 'withdraw')
        # همان پول نباید دوباره قابل درخواست باشد
        self.assertEqual(get_withdrawable_amount(self.user), 2_000_000)

    def test_cannot_request_more_than_balance(self):
        with self.assertRaises(ValueError):
            WithdrawalRequest.create_for(self.user, 9_000_000, VALID_IBAN, 'کسری ن')
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 5_000_000)
        # ردیف درخواست هم نباید باقی مانده باشد (اتمیک بودن create_for)
        self.assertEqual(WithdrawalRequest.objects.count(), 0)

    def test_reject_returns_the_money(self):
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        ok, err = req.reject(self.admin, reason='نام صاحب حساب مغایر است')
        self.assertTrue(ok, err)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 5_000_000)
        self.assertEqual(req.status, 'rejected')
        self.assertEqual(get_withdrawable_amount(self.user), 5_000_000)

    def test_reject_reference_does_not_pollute_the_payment_audit(self):
        """بازگشتِ ردِ درخواست نباید به‌عنوان «جبران پرداخت» شمرده شود.

        دستور audit_payment_ticket_balance هر تراکنشی با این پیشوندها را
        «پولی که به کاربر برگشته» حساب می‌کند. اگر بازگشتِ ردِ درخواستِ برداشت
        هم همان پیشوند را می‌گرفت، یک بدهیِ واقعیِ باشگاه در گزارش پنهان
        می‌شد -- کاربر پولش را در کیف پول دارد، نه از باشگاه طلبکار است.
        """
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        req.reject(self.admin, reason='تست')
        ref = Transaction.objects.filter(amount__gt=0).exclude(
            reference_id='compensate-64-9').first().reference_id
        self.assertEqual(ref, f'WD-REJECT-{req.pk}')
        self.assertFalse(ref.startswith(WITHDRAWABLE_PREFIXES))

    def test_paid_keeps_the_money_out(self):
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        req.approve(self.admin)
        self.assertEqual(req.status, 'approved')
        ok, err = req.mark_paid(self.admin, bank_reference='PAYA-123456')
        self.assertTrue(ok, err)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 2_000_000)
        self.assertEqual(req.bank_reference, 'PAYA-123456')

    def test_a_settled_request_cannot_be_processed_twice(self):
        """دو کلیک هم‌زمان ادمین نباید پول را دو بار برگرداند."""
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        req.reject(self.admin, reason='یک بار')
        ok, err = req.reject(self.admin, reason='دو بار')
        self.assertFalse(ok)
        self.assertIn('قبلاً', err)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 5_000_000)  # نه ۸ میلیون

    def test_paid_request_cannot_be_rejected_afterwards(self):
        req = WithdrawalRequest.create_for(self.user, 3_000_000, VALID_IBAN, 'کسری ن')
        req.mark_paid(self.admin, bank_reference='PAYA-1')
        ok, _ = req.reject(self.admin, reason='اشتباهی')
        self.assertFalse(ok)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 2_000_000)


class WithdrawalViewTests(TestCase):
    """ورودی کاربر: هیچ مسیری نباید بدون کلید فعال یا با ورودی نامعتبر رد شود."""

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u3', password='pw12345',
                                             user_type='normal')
        # کیف پول با سیگنال post_save روی User ساخته می‌شود، پس فقط برش می‌داریم.
        self.wallet = Wallet.objects.get(user=self.user)
        self.wallet.add_balance(amount=5_000_000, reference_id='compensate-64-3')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        # GeoAccessMiddleware هدر X-Iran-IP را از nginx انتظار دارد؛ تست‌کلاینت
        # آن را نمی‌فرستد و همه‌ی درخواست‌ها ۴۰۳ می‌شدند. موضوع این تست‌ها
        # کنترل جغرافیایی نیست.
        s.block_foreign_ips = False
        s.save()
        self.client.force_login(self.user)

    def _post(self, **kw):
        data = {'amount': '1000000', 'iban': VALID_IBAN,
                'account_holder': 'کسری نیک‌شبانی', 'national_code': '1234567890'}
        data.update(kw)
        return self.client.post('/wallet/withdraw/', data)

    def test_happy_path(self):
        resp = self._post()
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(WithdrawalRequest.objects.count(), 1)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 4_000_000)

    def test_invalid_iban_is_refused(self):
        self._post(iban='IR270170000000100324200000')
        self.assertEqual(WithdrawalRequest.objects.count(), 0)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.balance, 5_000_000)

    def test_amount_above_withdrawable_is_refused(self):
        self._post(amount='9000000')
        self.assertEqual(WithdrawalRequest.objects.count(), 0)

    def test_persian_digits_in_amount_and_iban_are_accepted(self):
        fa = str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹')
        self._post(amount='1000000'.translate(fa), iban=VALID_IBAN.translate(fa))
        self.assertEqual(WithdrawalRequest.objects.count(), 1)

    def test_second_open_request_is_refused(self):
        self._post()
        self._post()
        self.assertEqual(WithdrawalRequest.objects.count(), 1)

    def test_disabled_switch_closes_the_page(self):
        from accounts.models import SiteSettings
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = False
        s.save()
        resp = self._post()
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(WithdrawalRequest.objects.count(), 0)

    def test_vip_user_has_no_access(self):
        self.user.user_type = 'vip'
        self.user.save()
        self._post()
        self.assertEqual(WithdrawalRequest.objects.count(), 0)


class UserVisibilityAndCorrectionTests(TestCase):
    """آنچه کاربر می‌بیند و راهی که برای اصلاح اشتباهش دارد."""

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u9', password='pw12345')
        self.admin = User.objects.create_user(username='adm6', password='pw12345',
                                              is_staff=True, is_superuser=True)
        self.wallet = Wallet.objects.get(user=self.user)
        self.wallet.add_balance(amount=4_000_000, reference_id='compensate-64-90')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()
        self.client.force_login(self.user)
        self.req = WithdrawalRequest.create_for(
            self.user, 3_000_000, VALID_IBAN, 'کسری نیک‌شبانی')

    def test_user_sees_how_much_is_withdrawable(self):
        """کاربر باید بدون رفتن به فرم، عددِ قابل برداشت را ببیند."""
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')   # درخواست باز نباشد
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('قابل برداشت به حساب بانکی', body)
        self.assertIn('4٬000٬000', body)          # کل مبلغ جبرانی

    def test_held_amount_is_explained_not_silently_missing(self):
        """وقتی درخواستی باز است، «قابل برداشت» کمتر می‌شود -- باید توضیح بدهیم."""
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('نگه داشته شده', body)
        self.assertIn('3٬000٬000', body)          # مبلغِ نگه‌داشته‌شده

    def test_zero_withdrawable_is_explained(self):
        """کاربری که فقط شارژ شخصی دارد باید بفهمد چرا صفر است."""
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        other = User.objects.create_user(username='u12', password='pw12345')
        Wallet.objects.get(user=other).add_balance(
            amount=3_000_000, reference_id='4734690099')   # شارژ از درگاه
        self.client.force_login(other)
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('چیزی قابل برداشت نیست', body)
        self.assertIn('خودتان شارژ کرده‌اید', body)

    def test_nothing_is_shown_while_withdrawal_is_disabled(self):
        """تا وقتی کلید خاموش است، عددی که نمی‌شود رویش اقدام کرد نشان داده نشود."""
        from accounts.models import SiteSettings
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = False
        s.save()
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertNotIn('قابل برداشت به حساب بانکی', body)

    def test_bank_reference_reaches_the_user_wallet(self):
        self.req.mark_paid(self.admin, bank_reference='PAYA-556677')
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('PAYA-556677', body)
        self.assertIn('شناسه پیگیری واریز', body)

    def test_bank_reference_also_lands_in_the_transaction_history(self):
        """کاربر در تاریخچه‌ی تراکنش‌ها هم باید شماره‌ی پیگیری را ببیند."""
        self.req.mark_paid(self.admin, bank_reference='PAYA-998877')
        tx = Transaction.objects.get(user=self.user, reference_id=f'WD-{self.req.pk}')
        self.assertIn('PAYA-998877', tx.description)

    def test_rejection_reason_is_shown_to_the_user(self):
        self.req.reject(self.admin, reason='نام صاحب حساب با نام شما یکی نیست')
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('نام صاحب حساب با نام شما یکی نیست', body)
        self.assertIn('دلیل رد', body)

    def test_user_can_cancel_a_pending_request_and_get_the_money_back(self):
        resp = self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.assertEqual(resp.status_code, 302)
        self.req.refresh_from_db()
        self.wallet.refresh_from_db()
        self.assertEqual(self.req.status, 'cancelled')
        self.assertEqual(self.wallet.balance, 4_000_000)
        self.assertEqual(get_withdrawable_amount(self.user), 4_000_000)

    def test_after_cancelling_a_new_request_is_possible(self):
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.client.post('/wallet/withdraw/', {
            'amount': '2000000', 'iban': VALID_IBAN, 'account_holder': 'کسری نیک‌شبانی',
            'national_code': '1234567890'})
        self.assertEqual(WithdrawalRequest.objects.filter(
            user=self.user, status='pending').count(), 1)

    def test_an_approved_request_can_no_longer_be_cancelled_by_the_user(self):
        """خزانه‌داری ممکن است همین حالا در حال واریز باشد."""
        self.req.approve(self.admin)
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'approved')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 1_000_000)

    def test_a_user_cannot_cancel_someone_elses_request(self):
        other = User.objects.create_user(username='u10', password='pw12345')
        self.client.force_login(other)
        resp = self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.assertEqual(resp.status_code, 404)
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'pending')

    def test_cancelling_twice_does_not_return_the_money_twice(self):
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 4_000_000)

    def test_the_guide_is_on_the_withdraw_page(self):
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        body = self.client.get('/wallet/withdraw/').content.decode()
        self.assertIn('برداشت در سه مرحله', body)
        self.assertIn('شماره شبا را از بانک خودتان بگیرید', body)
        self.assertIn('اگر شبا را اشتباه وارد کنم', body)


class StatusFlowTests(TestCase):
    """چرخه‌ی وضعیت: در انتظار تأیید -> در انتظار پرداخت -> پرداخت انجام شد،
    و انشعابِ «نیاز به اصلاح اطلاعات»."""

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u11', password='pw12345')
        self.admin = User.objects.create_user(username='adm7', password='pw12345',
                                              is_staff=True, is_superuser=True)
        self.wallet = Wallet.objects.get(user=self.user)
        self.wallet.add_balance(amount=5_000_000, reference_id='compensate-64-95')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()
        self.client.force_login(self.user)
        self.req = WithdrawalRequest.create_for(
            self.user, 3_000_000, VALID_IBAN, 'کسری نیک‌شبانی')

    OTHER_IBAN = 'IR260620000000000123456789'

    # ---------------- در انتظار تأیید: قابل ویرایش ----------------
    def test_pending_request_is_editable_by_the_user(self):
        resp = self.client.post('/wallet/withdraw/', {
            'amount': '3000000', 'iban': self.OTHER_IBAN, 'account_holder': 'نام اصلاح‌شده',
            'national_code': '1234567890'})
        self.assertEqual(resp.status_code, 302)
        self.req.refresh_from_db()
        self.assertEqual(self.req.iban, self.OTHER_IBAN)
        self.assertEqual(self.req.account_holder, 'نام اصلاح‌شده')
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 1)

    def test_editing_the_amount_settles_the_difference(self):
        self.client.post('/wallet/withdraw/', {
            'amount': '2000000', 'iban': VALID_IBAN, 'account_holder': 'کسری نیک‌شبانی',
            'national_code': '1234567890'})
        self.req.refresh_from_db()
        self.wallet.refresh_from_db()
        self.assertEqual(self.req.amount, 2_000_000)
        self.assertEqual(self.wallet.balance, 3_000_000)      # ۱ میلیون برگشت
        self.assertEqual(get_withdrawable_amount(self.user), 3_000_000)

    def test_editing_upward_takes_the_difference(self):
        self.client.post('/wallet/withdraw/', {
            'amount': '4500000', 'iban': VALID_IBAN, 'account_holder': 'کسری نیک‌شبانی',
            'national_code': '1234567890'})
        self.req.refresh_from_db()
        self.wallet.refresh_from_db()
        self.assertEqual(self.req.amount, 4_500_000)
        self.assertEqual(self.wallet.balance, 500_000)

    def test_cannot_edit_above_what_is_withdrawable(self):
        self.client.post('/wallet/withdraw/', {
            'amount': '9000000', 'iban': VALID_IBAN, 'account_holder': 'کسری نیک‌شبانی',
            'national_code': '1234567890'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.amount, 3_000_000)

    def test_the_form_comes_prefilled_when_editing(self):
        body = self.client.get('/wallet/withdraw/').content.decode()
        self.assertIn(VALID_IBAN, body)
        self.assertIn('کسری نیک‌شبانی', body)

    # ---------------- در انتظار پرداخت: قفل ----------------
    def test_approved_request_is_locked_for_the_user(self):
        self.req.approve(self.admin)
        self.assertEqual(self.req.get_status_display(), 'در انتظار پرداخت')
        self.client.post('/wallet/withdraw/', {
            'amount': '1000000', 'iban': self.OTHER_IBAN, 'account_holder': 'نام دیگر',
            'national_code': '1234567890'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.iban, VALID_IBAN)
        self.assertEqual(self.req.amount, 3_000_000)

    def test_approved_request_cannot_be_cancelled_by_the_user(self):
        self.req.approve(self.admin)
        self.client.post(f'/wallet/withdraw/{self.req.pk}/cancel/')
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'approved')

    # ---------------- نیاز به اصلاح ----------------
    def test_admin_sends_it_back_and_the_money_stays_held(self):
        ok, err = self.req.request_correction(self.admin, reason='شبا با نام شما یکی نیست')
        self.assertTrue(ok, err)
        self.req.refresh_from_db()
        self.wallet.refresh_from_db()
        self.assertEqual(self.req.status, 'needs_correction')
        self.assertEqual(self.wallet.balance, 2_000_000, 'پول نباید برگشته باشد')
        self.assertEqual(get_withdrawable_amount(self.user), 2_000_000)

    def test_user_sees_what_to_fix_and_can_fix_it(self):
        self.req.request_correction(self.admin, reason='شبا با نام شما یکی نیست')
        body = self.client.get('/wallet/withdraw/').content.decode()
        self.assertIn('اطلاعات این درخواست باید اصلاح شود', body)
        self.assertIn('شبا با نام شما یکی نیست', body)

        self.client.post('/wallet/withdraw/', {
            'amount': '3000000', 'iban': self.OTHER_IBAN, 'account_holder': 'نام درست',
            'national_code': '1234567890'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'pending', 'بعد از اصلاح باید به صف بررسی برگردد')
        self.assertEqual(self.req.iban, self.OTHER_IBAN)
        self.assertEqual(self.req.admin_note, '', 'یادداشت قبلی باید پاک شود')

    def test_correction_can_also_come_from_the_paying_stage(self):
        """اگر بانک واریز را برگرداند، همان درخواست باید برای اصلاح برود."""
        self.req.approve(self.admin)
        ok, _ = self.req.request_correction(self.admin, reason='بانک شبا را نپذیرفت')
        self.assertTrue(ok)
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'needs_correction')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 2_000_000)

    def test_correction_route_through_the_admin_view(self):
        self.client.force_login(self.admin)
        self.client.post(f'/wallet/admin/withdrawals/{self.req.pk}/action/',
                         {'action': 'correct', 'note': 'کد ملی صاحب حساب مغایر است'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'needs_correction')
        self.assertEqual(self.req.admin_note, 'کد ملی صاحب حساب مغایر است')

    def test_correction_needs_a_reason(self):
        self.client.force_login(self.admin)
        self.client.post(f'/wallet/admin/withdrawals/{self.req.pk}/action/',
                         {'action': 'correct', 'note': ''})
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'pending')

    def test_a_second_request_is_still_blocked_while_correcting(self):
        """درخواست در حال اصلاح هنوز باز است -- نباید دومی ساخته شود."""
        self.req.request_correction(self.admin, reason='اصلاح کن')
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 1)
        self.client.post('/wallet/withdraw/', {
            'amount': '1000000', 'iban': self.OTHER_IBAN, 'account_holder': 'نام',
            'national_code': '1234567890'})
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 1)

    # ---------------- پرداخت انجام شد ----------------
    def test_full_happy_path(self):
        self.req.approve(self.admin)
        self.req.mark_paid(self.admin, bank_reference='PAYA-321')
        self.req.refresh_from_db()
        self.assertEqual(self.req.get_status_display(), 'پرداخت انجام شد')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 2_000_000)
        body = self.client.get('/wallet/dashboard/').content.decode()
        self.assertIn('PAYA-321', body)
        self.assertIn('پرداخت انجام شد', body)


class IbanInquiryTests(TestCase):
    """استعلام نام صاحب حساب از بانک.

    اصل حاکم: خرابیِ این سرویس هرگز نباید پول کسی را گیر بیندازد. فقط وقتی
    بانک *با اطمینان* نام دیگری برمی‌گرداند، مسیر عوض می‌شود.
    """

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u13', password='pw12345',
                                             first_name='کسری', last_name='نیک‌شبانی')
        self.admin = User.objects.create_user(username='adm8', password='pw12345',
                                              is_staff=True, is_superuser=True)
        self.wallet = Wallet.objects.get(user=self.user)
        self.wallet.add_balance(amount=5_000_000, reference_id='compensate-64-96')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()
        self.client.force_login(self.user)

    def _submit(self, **over):
        data = {'amount': '3000000', 'iban': VALID_IBAN,
                'account_holder': 'کسری نیک‌شبانی',
                'national_code': self.user.national_code or '1234567890'}
        data.update(over)
        return self.client.post('/wallet/withdraw/', data)

    @staticmethod
    def _zibal(result, data=None, message='موفق'):
        class R:
            status_code = 200
            @staticmethod
            def json():
                out = {'result': result, 'message': message}
                if data is not None:
                    out['data'] = data
                return out
        return R()

    # ---------------- تطابق ----------------
    def test_matching_name_marks_the_request_verified(self):
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(1, {'name': 'کسری نیک شبانی'})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertTrue(req.iban_verified)
        self.assertEqual(req.status, 'pending')
        self.assertEqual(req.iban_owner_name, 'کسری نیک شبانی')

    def test_spelling_variants_still_match(self):
        """«نیک‌شبانی» با نیم‌فاصله و «نيك شباني» عربی، همان نام‌اند."""
        from .iban_inquiry import names_match
        self.assertTrue(names_match('كسري نيك شباني', 'کسری نیک‌شبانی'))
        self.assertTrue(names_match('کسری نیک شبانی', 'کسری نیک‌شبانی'))
        self.assertTrue(names_match('کسری نیک‌شبانی', 'کسری نیک شبانی فرزند x'))
        self.assertFalse(names_match('مهدی احمدی', 'کسری نیک‌شبانی'))

    # ---------------- مغایرت ----------------
    def test_different_name_sends_it_to_correction_and_keeps_the_money(self):
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(1, {'name': 'مهدی احمدی'})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.wallet.refresh_from_db()
        self.assertFalse(req.iban_verified)
        self.assertEqual(req.status, 'needs_correction')
        self.assertIn('مهدی احمدی', req.admin_note)
        self.assertEqual(self.wallet.balance, 2_000_000, 'پول نباید برگشته باشد')

    def test_unknown_iban_at_the_bank_asks_for_correction(self):
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(21, message='شبا نامعتبر')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'needs_correction')
        self.assertFalse(req.iban_verified)

    # ---------------- سرویس در دسترس نیست ----------------
    def test_no_token_behaves_exactly_like_before(self):
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_zibal_wallet_empty_does_not_block_the_user(self):
        """کد ۲۹ یعنی کیف پول کارمزدِ *ما* خالی است -- تقصیر کاربر نیست."""
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(29, message='موجودی کافی نیست')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_iban_not_found_asks_the_user_to_fix_it(self):
        """کد ۴۴ («شبای مورد نظر یافت نشد») درباره‌ی ورودیِ کاربر است."""
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(44, message='شبا یافت نشد')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'needs_correction')
        self.assertFalse(req.iban_verified)

    def test_no_provider_available_is_unknown_not_mismatch(self):
        """کد ۴۵ درباره‌ی زیرساختِ استعلام است، نه شبای کاربر."""
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post',
                   return_value=self._zibal(45, message='سرویس دهنده ای در دسترس نیست')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_network_failure_does_not_block_the_user(self):
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=RuntimeError('timeout')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_unexpected_response_shape_is_treated_as_unknown_not_mismatch(self):
        """اگر شکل پاسخ عوض شود، نباید درخواست‌های درست را مغایر اعلام کنیم."""
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(1, {'somethingElse': 'x'})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_name_shapes_from_different_response_formats(self):
        from .iban_inquiry import _extract_name
        self.assertEqual(_extract_name({'name': 'الف ب'}), 'الف ب')
        self.assertEqual(_extract_name({'firstName': 'الف', 'lastName': 'ب'}), 'الف ب')
        self.assertEqual(_extract_name({'depositOwners': [{'fullName': 'الف ب'}]}), 'الف ب')
        self.assertEqual(_extract_name({}), '')

    # ---------------- تطابق کد ملی: حرفِ آخر ----------------
    def _routed(self, url_map):
        """پاسخ جعلی بر اساس آدرسی که صدا زده می‌شود."""
        def fake(url, **kw):
            for fragment, response in url_map.items():
                if fragment in url:
                    return response
            raise AssertionError(f'آدرس پیش‌بینی‌نشده: {url}')
        return fake

    def test_national_code_match_verifies_without_needing_the_name(self):
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': True})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertTrue(req.iban_verified)
        # تأیید قطعیِ بانک یعنی دیگر منتظر اپراتور نمی‌ماند
        self.assertEqual(req.status, 'approved')
        self.assertEqual(req.get_status_display(), 'در انتظار پرداخت')

    def test_happy_path_costs_only_one_inquiry(self):
        """در مسیر عادی نباید دو بار کارمزد بدهیم."""
        self.user.national_code = '1234567890'
        self.user.save()
        calls = []

        def fake(url, **kw):
            calls.append(url)
            return self._zibal(1, {'matched': True})
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), patch('requests.post', side_effect=fake):
            self._submit()
        self.assertEqual(len(calls), 1, f'بیش از یک استعلام زده شد: {calls}')

    def test_national_code_mismatch_beats_a_matching_name(self):
        """حتی اگر نام بخواند، وقتی کد ملی نمی‌خواند یعنی حساب مالِ او نیست.

        دو نفر می‌توانند هم‌نام باشند؛ این دقیقاً همان حالتی است که مقایسه‌ی
        نام از پسش برنمی‌آید.
        """
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': False}),
                 'ibanInquiry': self._zibal(1, {'name': 'کسری نیک‌شبانی'})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertFalse(req.iban_verified)
        self.assertEqual(req.status, 'needs_correction')
        self.assertIn('متعلق به شما نیست', req.admin_note)

    def test_mismatch_message_names_the_real_owner(self):
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': False}),
                 'ibanInquiry': self._zibal(1, {'name': 'مهدی احمدی'})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertIn('مهدی احمدی', req.admin_note)

    def test_falls_back_to_name_comparison_without_a_national_code(self):
        self.user.national_code = None
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'ibanInquiry': self._zibal(1, {'name': 'کسری نیک شبانی'})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertTrue(req.iban_verified)

    def test_match_service_down_falls_back_to_name(self):
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(45, message='در دسترس نیست'),
                 'ibanInquiry': self._zibal(1, {'name': 'کسری نیک شبانی'})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertTrue(req.iban_verified, 'با قطعیِ سرویس تطابق، باید به نام برگردد')

    def test_unknown_match_shape_is_not_treated_as_mismatch(self):
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'unexpected': 'x'}),
                 'ibanInquiry': self._zibal(45, message='در دسترس نیست')})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')
        self.assertIsNone(req.iban_verified)

    def test_birth_date_is_taken_from_past_purchases_when_available(self):
        from .iban_inquiry import find_birth_date
        from matches.models import Match, Stadium
        from payments.models import Payment
        from django.utils import timezone as tz
        from datetime import timedelta
        self.user.national_code = '1234567890'
        self.user.save()
        st = Stadium.objects.create(name='ورزشگاه', capacity=10)
        m = Match.objects.create(home_team='الف', away_team='ب', stadium=st,
                                 date_time=tz.now() + timedelta(days=1))
        Payment.objects.create(user=self.user, purpose='ticket_purchase', match=m,
                               status='success', gateway_amount=1000,
                               buyer_info={'national_code_5': '1234567890',
                                           'tarikhe_tavallod_5': '1365/03/12'})
        self.assertEqual(find_birth_date(self.user), '1365/03/12')

    # ---------------- قفلِ کد ملی ----------------
    def test_someone_elses_national_code_is_refused(self):
        """وگرنه می‌شد کد ملیِ صاحبِ واقعیِ یک حساب را زد، استعلام را پاس کرد
        و پول را به حساب همان شخص فرستاد."""
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': True})})):
            self._submit(national_code='9999999999')
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 0)

    def test_national_code_must_be_ten_digits(self):
        self.user.national_code = None
        self.user.save()
        self._submit(national_code='123')
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 0)

    def test_persian_digits_in_national_code_are_accepted(self):
        self.user.national_code = '1234567890'
        self.user.save()
        fa = str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹')
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit(national_code='1234567890'.translate(fa))
        self.assertEqual(WithdrawalRequest.objects.filter(user=self.user).count(), 1)

    def test_the_declared_code_is_stored_on_the_request(self):
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.national_code, '1234567890')

    # ---------------- تأیید خودکار ----------------
    def test_unverified_request_still_waits_for_a_human(self):
        """اگر نتوانستیم استعلام کنیم، تأیید خودکار نباید اتفاق بیفتد."""
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(45, message='در دسترس نیست')):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'pending')

    def test_name_only_match_does_not_auto_approve(self):
        """تطابق نام حدسی است -- هم‌نام بودن دو نفر ممکن است. برای رفتن به صف
        پرداخت، تأیید قطعیِ کد ملی لازم است."""
        self.user.national_code = None
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'ibanInquiry': self._zibal(1, {'name': 'کسری نیک شبانی'})})):
            self._submit(national_code='1234567890')
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertTrue(req.iban_verified)
        self.assertEqual(req.status, 'pending', 'تطابق نام نباید خودکار تأیید کند')

    def test_correcting_a_request_clears_the_old_verification(self):
        """جوابِ استعلام به شبای قبلی مربوط بود؛ با عوض شدن شبا باید دوباره پرسیده شود."""
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': False}),
                 'ibanInquiry': self._zibal(1, {'name': 'مهدی احمدی'})})):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertEqual(req.status, 'needs_correction')

        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': True})})):
            self._submit(iban='IR260620000000000123456789')
        req.refresh_from_db()
        self.assertEqual(req.iban, 'IR260620000000000123456789')
        self.assertTrue(req.iban_verified)
        self.assertEqual(req.status, 'approved')

    # ---------------- دستورِ بررسیِ درخواست‌های قدیمی ----------------
    def test_sweep_command_verifies_older_requests(self):
        from io import StringIO
        from django.core.management import call_command
        self.user.national_code = '1234567890'
        self.user.save()
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit()                              # بدون استعلام ثبت می‌شود
        req = WithdrawalRequest.objects.get(user=self.user)
        self.assertIsNone(req.iban_verified)

        out = StringIO()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': True})})):
            call_command('verify_withdrawal_ibans', '--execute', stdout=out, stderr=out)
        req.refresh_from_db()
        self.assertTrue(req.iban_verified)
        self.assertEqual(req.status, 'approved')

    def test_sweep_command_is_read_only_without_execute(self):
        from io import StringIO
        from django.core.management import call_command
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit()
        req = WithdrawalRequest.objects.get(user=self.user)
        out = StringIO()
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', side_effect=self._routed({
                 'checkIbanWithNationalCode': self._zibal(1, {'matched': True})})):
            call_command('verify_withdrawal_ibans', stdout=out, stderr=out)
        req.refresh_from_db()
        self.assertIsNone(req.iban_verified, 'حالت آزمایشی چیزی را تغییر داد')

    # ---------------- نمایش در پنل ----------------
    def test_admin_sees_the_bank_answer(self):
        with self.settings(ZIBAL_FACILITY_TOKEN='x'), \
             patch('requests.post', return_value=self._zibal(1, {'name': 'مهدی احمدی'})):
            self._submit()
        self.client.force_login(self.admin)
        body = self.client.get('/wallet/admin/withdrawals/?status=all').content.decode()
        self.assertIn('بانک نام دیگری برگرداند', body)
        self.assertIn('مهدی احمدی', body)

    def test_admin_sees_that_no_inquiry_happened(self):
        with self.settings(ZIBAL_FACILITY_TOKEN=''):
            self._submit()
        self.client.force_login(self.admin)
        body = self.client.get('/wallet/admin/withdrawals/?status=all').content.decode()
        self.assertIn('استعلام بانکی انجام نشد', body)


class PageRenderTests(TestCase):
    """صفحه‌ها واقعاً رندر شوند.

    یک خطای نحوی در تمپلیت را نه manage.py check می‌گیرد و نه تست‌های منطقی؛
    فقط رندر شدن واقعی آن را نشان می‌دهد.
    """

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u4', password='pw12345')
        self.admin = User.objects.create_user(username='adm2', password='pw12345',
                                              is_staff=True, is_superuser=True)
        w = Wallet.objects.get(user=self.user)
        w.add_balance(amount=4_000_000, reference_id='compensate-64-4')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()

    def test_user_dashboard_and_withdraw_page_render(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/wallet/dashboard/').status_code, 200)
        resp = self.client.get('/wallet/withdraw/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'حداکثر مبلغ قابل برداشت')

    def test_admin_queue_renders_with_and_without_rows(self):
        self.client.force_login(self.admin)
        url = '/wallet/admin/withdrawals/'
        self.assertEqual(self.client.get(url).status_code, 200)

        WithdrawalRequest.create_for(self.user, 1_000_000, VALID_IBAN, 'کسری نیک‌شبانی')
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'استعلام بانکی انجام نشد')
        for st in ('pending', 'approved', 'paid', 'rejected', 'all'):
            self.assertEqual(self.client.get(f'{url}?status={st}').status_code, 200)

    def test_admin_actions_move_money_through_the_view(self):
        self.client.force_login(self.admin)
        req = WithdrawalRequest.create_for(self.user, 1_000_000, VALID_IBAN, 'کسری نیک‌شبانی')
        url = f'/wallet/admin/withdrawals/{req.pk}/action/'

        # ثبت واریز بدون شماره پیگیری نباید انجام شود
        self.client.post(url, {'action': 'paid'})
        req.refresh_from_db()
        self.assertEqual(req.status, 'pending')

        # رد بدون دلیل هم نباید انجام شود
        self.client.post(url, {'action': 'reject'})
        req.refresh_from_db()
        self.assertEqual(req.status, 'pending')

        self.client.post(url, {'action': 'approve'})
        req.refresh_from_db()
        self.assertEqual(req.status, 'approved')

        self.client.post(url, {'action': 'paid', 'bank_reference': 'PAYA-9'})
        req.refresh_from_db()
        self.assertEqual(req.status, 'paid')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 3_000_000)

    def test_site_settings_page_shows_the_withdrawal_card(self):
        """کارت «برداشت وجه» در تنظیمات سایت، عددِ واقعیِ قابل برداشت را نشان دهد."""
        self.client.force_login(self.admin)
        resp = self.client.get('/accounts/admin/site-settings/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'برداشت وجه از کیف پول')
        self.assertEqual(resp.context['withdrawable_total'], 4_000_000)
        self.assertEqual(resp.context['withdrawable_users'], 1)

    def test_non_staff_cannot_open_the_admin_queue(self):
        self.client.force_login(self.user)
        resp = self.client.get('/wallet/admin/withdrawals/')
        self.assertNotEqual(resp.status_code, 200)


class TemplateCommentLintTests(TestCase):
    """کامنت `{# ... #}` چندخطی نباید در هیچ تمپلیتی وجود داشته باشد.

    لکسر جنگو (django/template/base.py) با الگوی
    `({%.*?%}|{{.*?}}|{#.*?#})` و **بدون** فلگ DOTALL کار می‌کند، یعنی نقطه
    خط جدید را نمی‌گیرد. پس `{#` که تا خط بعد ادامه پیدا کند اصلاً کامنت
    شناخته نمی‌شود و متنش خام به کاربر نمایش داده می‌شود. این دقیقاً روی
    صفحه‌ی کیف پول در پروداکشن اتفاق افتاد.

    برای توضیح چندخطی باید از {% comment %}...{% endcomment %} استفاده شود.
    """

    def test_no_multiline_hash_comments_in_templates(self):
        from pathlib import Path
        from django.conf import settings

        # همان رفتار لکسر تقلید می‌شود: از چپ به راست، و هر `{#` که بسته شد،
        # مکان‌نما بعد از `#}` می‌پرد. بدون این پرش، `{#` هایی که *داخل* یک
        # کامنت بسته‌شده‌اند (مثل `{# setTimeout(function () {#}`) اشتباهاً
        # خطا گزارش می‌شوند -- همان‌طور که جنگو هم آن‌ها را نمی‌بیند.
        offenders = []
        for tpl_dir in settings.TEMPLATES[0]['DIRS']:
            for path in sorted(Path(tpl_dir).rglob('*.html')):
                for lineno, line in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
                    pos = 0
                    while (start := line.find('{#', pos)) != -1:
                        close = line.find('#}', start + 2)
                        if close == -1:
                            offenders.append(f'{path}:{lineno}  {line.strip()[:80]}')
                            break
                        pos = close + 2
        self.assertEqual(offenders, [], 'کامنت {# #} چندخطی پیدا شد (خام رندر می‌شود):\n'
                                        + '\n'.join(offenders))


class RenderedPagesHaveNoLeakedCommentsTests(TestCase):
    """هیچ متنِ توضیحیِ داخل کد نباید در خروجی صفحه دیده شود."""

    def setUp(self):
        from accounts.models import SiteSettings
        self.user = User.objects.create_user(username='u5', password='pw12345')
        self.admin = User.objects.create_user(username='adm3', password='pw12345',
                                              is_staff=True, is_superuser=True)
        Wallet.objects.get(user=self.user).add_balance(
            amount=4_000_000, reference_id='compensate-64-5')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()

    def test_wallet_pages_do_not_show_source_comments(self):
        self.client.force_login(self.user)
        for url in ['/wallet/dashboard/', '/wallet/withdraw/']:
            body = self.client.get(url).content.decode()
            self.assertNotIn('{#', body, url)
            self.assertNotIn('wallet_charge_enabled', body, url)

    def test_admin_pages_do_not_show_source_comments(self):
        self.client.force_login(self.admin)
        WithdrawalRequest.create_for(self.user, 1_000_000, VALID_IBAN, 'کسری نیک‌شبانی')
        for url in ['/wallet/admin/withdrawals/', '/accounts/admin/site-settings/']:
            body = self.client.get(url).content.decode()
            self.assertNotIn('{#', body, url)
            self.assertNotIn('سیستم شبا را استعلام', body, url)


class FinanceFileTests(TestCase):
    """چرخه‌ی خزانه‌داری: فایل بگیر، پرداخت کن، همان فایل را برگردان."""

    def setUp(self):
        from accounts.models import SiteSettings
        self.admin = User.objects.create_user(username='adm4', password='pw12345',
                                              is_staff=True, is_superuser=True)
        self.user = User.objects.create_user(username='u6', password='pw12345',
                                             phone_number='09120001111',
                                             national_code='1111111111')
        Wallet.objects.get(user=self.user).add_balance(
            amount=5_000_000, reference_id='compensate-64-6')
        s = SiteSettings.get_solo()
        s.withdrawal_enabled = True
        s.block_foreign_ips = False
        s.save()
        self.req = WithdrawalRequest.create_for(
            self.user, 3_000_000, VALID_IBAN, 'کسری نیک‌شبانی')
        self.req.approve(self.admin)
        self.client.force_login(self.admin)

    # ------------------------------------------------------------------
    def _download(self, fmt=''):
        url = '/wallet/admin/withdrawals/export/?status=approved'
        if fmt:
            url += f'&format={fmt}'
        return self.client.get(url)

    def test_excel_export_carries_what_the_bank_needs(self):
        from io import BytesIO
        from openpyxl import load_workbook
        resp = self._download()
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(BytesIO(resp.content)).active
        header = [c.value for c in ws[1]]
        self.assertIn('شماره شبا', header)
        self.assertIn('مبلغ (ریال)', header)
        self.assertIn('شماره پیگیری واریز', header)
        row = [c.value for c in ws[2]]
        self.assertEqual(row[0], self.req.id)
        self.assertEqual(row[2], VALID_IBAN)
        self.assertEqual(row[3], 3_000_000)
        self.assertIn('', [row[8] or ''])          # ستون پیگیری خالی است

    def test_csv_export_opens_correctly_in_excel(self):
        resp = self._download('csv')
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8-sig')
        self.assertIn('شماره شبا', body)
        self.assertIn(VALID_IBAN, body)
        # بدون BOM اکسل فارسی ستون‌ها را به‌هم می‌ریزد
        self.assertTrue(resp.content.startswith('﻿'.encode('utf-8')))

    def test_export_only_includes_approved_by_default(self):
        other = User.objects.create_user(username='u7', password='x')
        Wallet.objects.get(user=other).add_balance(
            amount=2_000_000, reference_id='compensate-64-77')
        WithdrawalRequest.create_for(other, 1_000_000, VALID_IBAN, 'دیگری')  # pending
        body = self._download('csv').content.decode('utf-8-sig')
        self.assertEqual(body.count('IR27'), 1, 'درخواستِ بررسی‌نشده هم در فایل آمد')

    # ------------------------------------------------------------------
    def _upload(self, rows):
        import csv
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(['شماره درخواست', 'نام', 'شبا', 'مبلغ', 'نام کاربر',
                    'موبایل', 'کد ملی', 'تاریخ', 'شماره پیگیری واریز'])
        for r in rows:
            w.writerow(r)
        f = SimpleUploadedFile('paid.csv', buf.getvalue().encode('utf-8-sig'),
                               content_type='text/csv')
        return self.client.post('/wallet/admin/withdrawals/import/', {'file': f})

    def test_uploading_the_filled_file_marks_them_paid(self):
        self._upload([[self.req.id, 'x', VALID_IBAN, 3_000_000, 'x', '', '', '', 'PAYA-77']])
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'paid')
        self.assertEqual(self.req.bank_reference, 'PAYA-77')

    def test_row_without_a_bank_reference_is_left_alone(self):
        self._upload([[self.req.id, 'x', VALID_IBAN, 3_000_000, 'x', '', '', '', '']])
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'approved')

    def test_a_changed_amount_is_refused(self):
        """اگر مبلغ فایل با درخواست یکی نباشد یعنی مبلغ دیگری پرداخت شده."""
        self._upload([[self.req.id, 'x', VALID_IBAN, 9_000_000, 'x', '', '', '', 'PAYA-9']])
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'approved', 'با مبلغ مغایر ثبت شد')

    def test_uploading_twice_does_not_double_process(self):
        self._upload([[self.req.id, 'x', VALID_IBAN, 3_000_000, 'x', '', '', '', 'PAYA-1']])
        self._upload([[self.req.id, 'x', VALID_IBAN, 3_000_000, 'x', '', '', '', 'PAYA-2']])
        self.req.refresh_from_db()
        self.assertEqual(self.req.bank_reference, 'PAYA-1', 'ثبت دوباره وضعیت را عوض کرد')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 2_000_000)

    def test_persian_digits_in_the_returned_file_are_understood(self):
        fa = str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹')
        self._upload([[str(self.req.id).translate(fa), 'x', VALID_IBAN,
                       '3,000,000'.translate(fa), 'x', '', '', '', 'PAYA-FA']])
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, 'paid')

    def test_non_staff_cannot_download_the_file(self):
        self.client.force_login(self.user)
        self.assertNotEqual(self._download().status_code, 200)


class AdminCreditTests(TestCase):
    """پولی که باشگاه از پنل واریز می‌کند باید قابل برداشت باشد.

    بدون این صفحه، تنها راهِ در دسترسِ ادمین ویرایش مستقیم موجودی در پنل
    جنگو بود -- کاری که عدد را عوض می‌کند ولی هیچ تراکنشی نمی‌سازد، پس آن
    پول نه در تاریخچه دیده می‌شود، نه ممیزی می‌بیندش، و نه قابل برداشت
    می‌شود. یعنی دقیقاً همان چیزی که باشگاه می‌خواست، کار نمی‌کرد.
    """

    def setUp(self):
        from accounts.models import SiteSettings
        self.admin = User.objects.create_user(username='adm5', password='pw12345',
                                              is_staff=True, is_superuser=True)
        self.user = User.objects.create_user(username='u8', password='pw12345',
                                             phone_number='09121234567',
                                             national_code='2222222222')
        s = SiteSettings.get_solo()
        s.block_foreign_ips = False
        s.save()
        self.client.force_login(self.admin)

    def _credit(self, **over):
        data = {'action': 'credit', 'user_id': self.user.id, 'amount': '2000000',
                'reason': 'بلیط به دلیل خطای سامانه صادر نشد', 'q': '09121234567'}
        data.update(over)
        return self.client.post('/wallet/admin/wallet-credit/', data, follow=True)

    def test_credited_money_is_immediately_withdrawable(self):
        self._credit()
        w = Wallet.objects.get(user=self.user)
        self.assertEqual(w.balance, 2_000_000)
        self.assertEqual(get_withdrawable_amount(self.user), 2_000_000)

    def test_it_leaves_a_traceable_transaction(self):
        self._credit()
        tx = Transaction.objects.get(user=self.user, amount=2_000_000)
        self.assertIn('خطای سامانه', tx.description)
        self.assertIn(self.admin.username, tx.description)
        self.assertTrue(tx.reference_id.startswith('ADMIN-CREDIT-'))

    def test_attaching_a_match_makes_the_audit_see_it(self):
        from datetime import timedelta
        from django.utils import timezone as tz
        from matches.models import Match, Stadium
        st = Stadium.objects.create(name='ورزشگاه', capacity=100)
        m = Match.objects.create(home_team='الف', away_team='ب', stadium=st,
                                 date_time=tz.now() + timedelta(days=1))
        self._credit(match_id=str(m.id))
        tx = Transaction.objects.get(user=self.user, amount=2_000_000)
        self.assertTrue(tx.reference_id.startswith(f'compensate-{m.id}-'))
        self.assertEqual(get_withdrawable_amount(self.user), 2_000_000)

    def test_reason_is_required(self):
        self._credit(reason='کم')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 0)

    def test_zero_or_negative_is_refused(self):
        self._credit(amount='0')
        self._credit(amount='-500000')
        self.assertEqual(Wallet.objects.get(user=self.user).balance, 0)

    def test_non_staff_cannot_open_it(self):
        self.client.force_login(self.user)
        r = self.client.get('/wallet/admin/wallet-credit/')
        self.assertNotEqual(r.status_code, 200)

    def test_cancelled_match_refund_is_withdrawable(self):
        """پولی که بابت لغو مسابقه برگشته هم مال باشگاه است، نه شارژ کاربر."""
        Wallet.objects.get(user=self.user).add_balance(
            amount=1_500_000, reference_id='CANCEL-MATCH-64', tx_type='refund')
        self.assertEqual(get_withdrawable_amount(self.user), 1_500_000)


class IbanValidationTests(TestCase):
    def test_accepts_common_input_shapes(self):
        body = VALID_IBAN[2:]
        spaced = 'IR ' + ' '.join(body[i:i + 4] for i in range(0, len(body), 4))
        for raw in [VALID_IBAN, VALID_IBAN.lower(), spaced, body,
                    VALID_IBAN.translate(str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹'))]:
            self.assertTrue(is_valid_iban(normalize_iban(raw)), raw)

    def test_rejects_wrong_checksum_and_foreign_iban(self):
        for raw in ['', 'IR123', 'IR000000000000000000000000',
                    'DE89370400440532013000', 'IR' + 'A' * 24]:
            self.assertFalse(is_valid_iban(normalize_iban(raw)), raw)
