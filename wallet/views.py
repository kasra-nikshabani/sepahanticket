# wallet/views.py
from django.core.paginator import Paginator, PageNotAnInteger, EmptyPage
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.views.decorators.cache import never_cache

from .models import (Wallet, Transaction, WithdrawalRequest,
                     is_wallet_enabled, is_wallet_charge_enabled, is_withdrawal_enabled,
                     get_withdrawable_amount, normalize_iban, is_valid_iban,
                     MIN_WITHDRAWAL_AMOUNT,
                     WALLET_DISABLED_MESSAGE, WALLET_CHARGE_DISABLED_MESSAGE,
                     WITHDRAWAL_DISABLED_MESSAGE)


@never_cache
@login_required
def wallet_dashboard(request):
    """نمایش کیف پول و تاریخچه تراکنش‌ها با صفحه‌بندی"""
    if request.user.user_type == 'vip':
        messages.error(request, 'کاربران ویژه به این بخش دسترسی ندارند.')
        return redirect('matches:home')

    wallet, created = Wallet.objects.get_or_create(user=request.user)

    transactions_list = Transaction.objects.filter(
        user=request.user
    ).order_by('-created_at')

    paginator = Paginator(transactions_list, 5)
    page = request.GET.get('page', 1)

    try:
        transactions = paginator.page(page)
    except PageNotAnInteger:
        transactions = paginator.page(1)
    except EmptyPage:
        transactions = paginator.page(paginator.num_pages)

    total_deposits = Transaction.objects.filter(
        user=request.user,
        transaction_type='deposit'
    ).aggregate(total=Sum('amount'))['total'] or 0

    total_withdrawals = Transaction.objects.filter(
        user=request.user,
        transaction_type__in=['withdraw', 'ticket_purchase']
    ).aggregate(total=Sum('amount'))['total'] or 0

    total_ticket_purchases = Transaction.objects.filter(
        user=request.user,
        transaction_type='ticket_purchase'
    ).aggregate(total=Sum('amount'))['total'] or 0

    total_refunds = Transaction.objects.filter(
        user=request.user,
        transaction_type='refund'
    ).aggregate(total=Sum('amount'))['total'] or 0

    total_transactions = Transaction.objects.filter(user=request.user).count()

    context = {
        'wallet': wallet,
        'wallet_balance': wallet.balance,
        'transactions': transactions,
        'total_deposits': total_deposits,
        'total_withdrawals': total_withdrawals,
        'total_ticket_purchases': total_ticket_purchases,
        'total_refunds': total_refunds,
        'total_transactions': total_transactions,
        # داشبورد در حالت غیرفعال هم باز می‌ماند (کاربر باید موجودی و تاریخچه‌اش
        # را ببیند)، ولی دکمه‌های شارژ جای خود را به یک توضیح می‌دهند.
        'wallet_enabled': is_wallet_enabled(),
        'wallet_charge_enabled': is_wallet_charge_enabled(),
        'withdrawal_enabled': is_withdrawal_enabled(),
        'withdrawable': get_withdrawable_amount(request.user),
        'withdrawals': WithdrawalRequest.objects.filter(user=request.user)[:10],
    }
    return render(request, 'wallet/dashboard.html', context)


@login_required
def wallet_charge(request):
    """
    صفحه‌ی شارژ کیف پول.

    نکته‌ی مهم: این ویو دیگر خودش موجودی را افزایش نمی‌دهد (قبلاً یک نسخه‌ی
    شبیه‌سازی‌شده بدون پرداخت واقعی بود که حذف شد). شارژ واقعی همیشه از طریق
    فرم این صفحه مستقیماً به payments:payment_request ارسال می‌شود و فقط پس
    از تایید واقعی زیبال در payments:payment_verify انجام می‌گیرد.
    """
    if request.user.user_type == 'vip':
        messages.error(request, 'کاربران ویژه به این بخش دسترسی ندارند.')
        return redirect('matches:home')

    # ===== کیف پول غیرفعال یعنی صفحه‌ی شارژ اصلاً باز نشود =====
    # پنهان کردن دکمه در تمپلیت کافی نیست؛ آدرس این صفحه ممکن است بوکمارک
    # شده باشد یا مستقیم تایپ شود.
    if not is_wallet_charge_enabled():
        messages.error(request, WALLET_CHARGE_DISABLED_MESSAGE)
        return redirect('wallet:dashboard')

    wallet = get_object_or_404(Wallet, user=request.user)

    if request.method == 'POST':
        # این مسیر دیگر پردازشی انجام نمی‌دهد؛ فرم صفحه باید مستقیم به
        # payments:payment_request ارسال شود (طبق تمپلیت فعلی charge.html).
        messages.info(request, 'برای شارژ کیف پول از دکمه‌ی «پرداخت و شارژ» استفاده کنید.')
        return redirect('wallet:charge')

    return render(request, 'wallet/charge.html', {'wallet': wallet})


# ============================================================
#  برداشت وجه -- سمت کاربر
# ============================================================

@never_cache
@login_required
def wallet_withdraw(request):
    """ثبت درخواست برداشت به شبا.

    این ویو پول را جابه‌جا نمی‌کند؛ فقط درخواست می‌سازد و موجودی را نگه
    می‌دارد. واریز واقعی را خزانه‌دار از پنل انجام می‌دهد.
    """
    if request.user.user_type == 'vip':
        messages.error(request, 'کاربران ویژه به این بخش دسترسی ندارند.')
        return redirect('matches:home')

    if not is_withdrawal_enabled():
        messages.error(request, WITHDRAWAL_DISABLED_MESSAGE)
        return redirect('wallet:dashboard')

    withdrawable = get_withdrawable_amount(request.user)
    pending = WithdrawalRequest.objects.filter(
        user=request.user, status__in=WithdrawalRequest.OPEN_STATUSES
    ).first()

    # مقادیری که اگر فرم رد شد، دوباره در آن نشان داده شوند تا کاربر مجبور
    # به تایپ دوباره‌ی شبا نباشد.
    form = {
        'amount': request.POST.get('amount', '') if request.method == 'POST' else '',
        'iban': request.POST.get('iban', '') if request.method == 'POST' else '',
        'account_holder': (request.POST.get('account_holder', '') if request.method == 'POST'
                           else (request.user.get_full_name() or '').strip()),
    }

    def _render(error=None):
        if error:
            messages.error(request, error)
        return render(request, 'wallet/withdraw.html', {
            'withdrawable': withdrawable,
            'pending': pending,
            'min_amount': MIN_WITHDRAWAL_AMOUNT,
            'form': form,
        })

    if request.method != 'POST':
        return _render()

    if pending:
        return _render('شما یک درخواست برداشت در جریان دارید؛ تا تعیین تکلیف آن '
                       'نمی‌توانید درخواست تازه‌ای ثبت کنید.')

    if withdrawable < MIN_WITHDRAWAL_AMOUNT:
        return _render(f'مبلغ قابل برداشت شما کمتر از حداقل مجاز '
                       f'({MIN_WITHDRAWAL_AMOUNT:,} ریال) است.')

    try:
        amount = int(str(request.POST.get('amount', '')).translate(
            str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
        ).replace(',', '').replace('٬', '').strip())
    except (TypeError, ValueError):
        return _render('مبلغ را به عدد وارد کنید.')

    if amount < MIN_WITHDRAWAL_AMOUNT:
        return _render(f'حداقل مبلغ برداشت {MIN_WITHDRAWAL_AMOUNT:,} ریال است.')
    if amount > withdrawable:
        return _render(f'حداکثر مبلغ قابل برداشت شما {withdrawable:,} ریال است.')

    iban = normalize_iban(request.POST.get('iban'))
    if not is_valid_iban(iban):
        return _render('شماره شبا معتبر نیست. شبا باید ۲۴ رقم بعد از IR داشته باشد.')

    holder = (request.POST.get('account_holder') or '').strip()
    if len(holder) < 3:
        return _render('نام صاحب حساب را کامل وارد کنید.')

    try:
        req = WithdrawalRequest.create_for(request.user, amount, iban, holder)
    except ValueError:
        # موجودی بین بارگذاری فرم و ارسال آن خرج شده است.
        return _render('موجودی کیف پول شما تغییر کرده است؛ لطفاً دوباره تلاش کنید.')

    messages.success(
        request,
        f'درخواست برداشت #{req.pk} به مبلغ {amount:,} ریال ثبت شد. '
        'پس از بررسی و تأیید، مبلغ به حساب اعلامی شما واریز می‌شود.'
    )
    return redirect('wallet:dashboard')


# ============================================================
#  برداشت وجه -- پنل مدیریت
# ============================================================

@never_cache
@staff_member_required
def admin_withdrawal_list(request):
    """صف بررسی درخواست‌های برداشت.

    ترتیب عمداً «قدیمی‌ترینِ در انتظار، اول» است: کسی که زودتر درخواست داده
    نباید پشت درخواست‌های تازه بماند.
    """
    status = request.GET.get('status', 'open')
    qs = WithdrawalRequest.objects.select_related('user', 'processed_by')

    if status == 'open':
        qs = qs.filter(status__in=WithdrawalRequest.OPEN_STATUSES).order_by('created_at')
    elif status in dict(WithdrawalRequest.STATUS_CHOICES):
        qs = qs.filter(status=status)
    # 'all' -> بدون فیلتر

    q = (request.GET.get('q') or '').strip()
    if q:
        from django.db.models import Q
        digits = q.translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789'))
        qs = qs.filter(
            Q(user__username__icontains=q) | Q(user__phone_number__icontains=digits)
            | Q(user__national_code__icontains=digits) | Q(iban__icontains=digits)
            | Q(account_holder__icontains=q)
        )

    counts = {
        'pending': WithdrawalRequest.objects.filter(status='pending').count(),
        'approved': WithdrawalRequest.objects.filter(status='approved').count(),
    }
    open_total = WithdrawalRequest.objects.filter(
        status__in=WithdrawalRequest.OPEN_STATUSES
    ).aggregate(s=Sum('amount'))['s'] or 0
    paid_total = WithdrawalRequest.objects.filter(
        status='paid'
    ).aggregate(s=Sum('amount'))['s'] or 0

    paginator = Paginator(qs, 50)
    try:
        page = paginator.page(request.GET.get('page', 1))
    except PageNotAnInteger:
        page = paginator.page(1)
    except EmptyPage:
        page = paginator.page(paginator.num_pages)

    return render(request, 'wallet/admin_withdrawals.html', {
        'requests': page,
        'status': status,
        'q': q,
        'counts': counts,
        'open_total': open_total,
        'paid_total': paid_total,
        'withdrawal_enabled': is_withdrawal_enabled(),
    })


@staff_member_required
def admin_withdrawal_action(request, request_id):
    """تأیید / ثبت واریز / رد یک درخواست."""
    if request.method != 'POST':
        return redirect('wallet:admin_withdrawal_list')

    wr = get_object_or_404(WithdrawalRequest, pk=request_id)
    action = request.POST.get('action')
    note = (request.POST.get('note') or '').strip()

    if action == 'approve':
        ok, err = wr.approve(request.user, note=note)
        done = f'درخواست #{wr.pk} تأیید شد؛ حالا در انتظار واریز است.'
    elif action == 'paid':
        bank_ref = (request.POST.get('bank_reference') or '').strip()
        if not bank_ref:
            messages.error(request, 'برای ثبت واریز، شماره پیگیری بانکی الزامی است.')
            return redirect('wallet:admin_withdrawal_list')
        ok, err = wr.mark_paid(request.user, bank_reference=bank_ref, note=note)
        done = f'واریز درخواست #{wr.pk} ثبت شد.'
    elif action == 'reject':
        if not note:
            messages.error(request, 'برای رد درخواست، ذکر دلیل الزامی است '
                                    '(همان متن برای کاربر نمایش داده می‌شود).')
            return redirect('wallet:admin_withdrawal_list')
        ok, err = wr.reject(request.user, reason=note)
        done = f'درخواست #{wr.pk} رد شد و مبلغ {wr.amount:,} ریال به کیف پول کاربر برگشت.'
    else:
        messages.error(request, 'درخواست نامعتبر است.')
        return redirect('wallet:admin_withdrawal_list')

    if ok:
        messages.success(request, done)
    else:
        messages.warning(request, err)

    # کاربر را به همان فیلتر/صفحه‌ای برمی‌گردانیم که از آن آمده بود، وگرنه
    # بعد از هر تأیید باید دوباره فیلترها را بچیند.
    back = (request.POST.get('back_query') or '').lstrip('?')
    url = reverse('wallet:admin_withdrawal_list')
    return redirect(f'{url}?{back}' if back else url)


@login_required
def wallet_withdraw_cancel(request, request_id):
    """کاربر درخواست خودش را پس می‌گیرد -- راهِ اصلاحِ اطلاعات اشتباه.

    بدون این، کسی که شبا را غلط زده باید منتظر بماند تا مدیر رد کند؛ در آن
    فاصله پولش هم بلوکه است و هم نمی‌تواند درخواست درست ثبت کند.
    """
    if request.method != 'POST':
        return redirect('wallet:dashboard')

    wr = get_object_or_404(WithdrawalRequest, pk=request_id, user=request.user)
    ok, err = wr.cancel_by_user()
    if ok:
        messages.success(
            request,
            f'درخواست #{wr.pk} لغو شد و {wr.amount:,} ریال به کیف پول شما برگشت. '
            'حالا می‌توانید درخواست تازه با اطلاعات درست ثبت کنید.')
    else:
        messages.error(request, err)
    return redirect('wallet:dashboard')


# ============================================================
#  شارژ جبرانیِ دستی از پنل
# ============================================================
# چرا لازم شد
# -----------
# تا امروز هیچ راهِ درستی برای «باشگاه بابت مشکل سامانه پول به کیف پول
# کاربر بریزد» از داخل پنل وجود نداشت. تنها راه‌ها دستورهای مدیریتی روی
# سرور بودند. تنها گزینه‌ی به‌ظاهر موجود -- ویرایش مستقیم موجودی در پنل
# جنگو -- عملاً خراب است: عدد را عوض می‌کند ولی هیچ تراکنشی نمی‌سازد، پس
# نه ردی در تاریخچه می‌ماند، نه ممیزی آن را می‌بیند، و نه آن پول قابل
# برداشت می‌شود (چون «قابل برداشت» از روی مرجعِ تراکنش تشخیص داده می‌شود).
# این ویو همان کار را درست انجام می‌دهد.

@never_cache
@staff_member_required
def admin_wallet_credit(request):
    from matches.models import Match
    from accounts.models import User

    matches = Match.objects.filter(is_active=True).order_by('-id')[:20]
    found = None
    query = (request.GET.get('q') or request.POST.get('q') or '').strip()
    if query:
        digits = query.translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789'))
        found = list(User.objects.filter(
            Q(phone_number__icontains=digits) | Q(national_code__icontains=digits)
            | Q(username__icontains=query)
        ).exclude(user_type='vip')[:10])

    if request.method == 'POST' and request.POST.get('action') == 'credit':
        user_id = request.POST.get('user_id')
        reason = (request.POST.get('reason') or '').strip()
        match_id = request.POST.get('match_id') or ''
        try:
            amount = int(str(request.POST.get('amount', '')).translate(
                str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
            ).replace(',', '').replace('٬', '').strip())
        except (TypeError, ValueError):
            amount = 0

        target = User.objects.filter(pk=user_id).first()
        if target is None:
            messages.error(request, 'کاربر پیدا نشد.')
        elif amount <= 0:
            messages.error(request, 'مبلغ باید بزرگ‌تر از صفر باشد.')
        elif len(reason) < 5:
            messages.error(request, 'دلیل شارژ را بنویسید — در تاریخچه‌ی کاربر ثبت می‌شود.')
        else:
            # مرجع تعیین می‌کند این پول «برگشتیِ باشگاه» شمرده شود و در نتیجه
            # قابل برداشت باشد. اگر به مسابقه‌ای وصل شود، ممیزیِ همان مسابقه
            # هم آن را به‌عنوان جبرانِ پرداخت حساب می‌کند.
            stamp = timezone.now().strftime('%Y%m%d%H%M%S')
            if match_id.isdigit():
                reference = f'compensate-{match_id}-manual{stamp}'
            else:
                reference = f'ADMIN-CREDIT-{stamp}-{target.id}'

            wallet, _ = Wallet.objects.get_or_create(user=target)
            ok = wallet.add_balance(
                amount=amount,
                description=f'{reason} (ثبت توسط {request.user.username})',
                reference_id=reference,
                tx_type='refund',
            )
            if ok:
                messages.success(
                    request,
                    f'{amount:,} ریال به کیف پول {target.phone_number or target.username} '
                    f'اضافه شد و قابل برداشت است.')
                return redirect(f"{reverse('wallet:admin_wallet_credit')}?q={query}")
            messages.error(request, 'واریز انجام نشد.')

    return render(request, 'wallet/admin_wallet_credit.html', {
        'q': query, 'found': found, 'matches': matches,
    })


# ============================================================
#  خروجی برای خزانه‌داری و برگرداندن نتیجه‌ی واریز
# ============================================================
# ستون‌ها عمداً همین ترتیب‌اند: شناسه‌ی درخواست اول می‌آید چون کلیدِ برگرداندن
# فایل است، و «شماره پیگیری واریز» آخرین ستونِ خالی است تا خزانه‌دار فقط همان
# را پر کند. ستون‌های نام/موبایل/کد ملی برای همان مقایسه‌ی چشمی‌اند که تنها
# کنترل مالکیت حساب است.
EXPORT_COLUMNS = [
    'شماره درخواست', 'نام صاحب حساب', 'شماره شبا', 'مبلغ (ریال)',
    'نام کاربر در سامانه', 'موبایل', 'کد ملی', 'تاریخ درخواست',
    'شماره پیگیری واریز',
]


def _export_rows(qs):
    for r in qs:
        yield [
            r.id,
            r.account_holder,
            r.iban,
            r.amount,
            (r.user.get_full_name() or r.user.username),
            r.user.phone_number or '',
            r.user.national_code or '',
            r.created_at.strftime('%Y/%m/%d %H:%M'),
            r.bank_reference or '',
        ]


@never_cache
@staff_member_required
def admin_withdrawal_export(request):
    """فایل پرداخت دسته‌ای برای خزانه‌داری (اکسل یا CSV).

    پیش‌فرض فقط درخواست‌های «تأیید شده -- در انتظار واریز» را می‌دهد؛ چیزی که
    هنوز بررسی نشده نباید سهواً پرداخت شود.
    """
    status = request.GET.get('status', 'approved')
    qs = WithdrawalRequest.objects.select_related('user')
    if status == 'open':
        qs = qs.filter(status__in=WithdrawalRequest.OPEN_STATUSES)
    elif status in dict(WithdrawalRequest.STATUS_CHOICES):
        qs = qs.filter(status=status)
    qs = qs.order_by('created_at')

    stamp = timezone.now().strftime('%Y%m%d-%H%M')
    rows = list(_export_rows(qs))

    if request.GET.get('format') == 'csv':
        import csv
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="withdrawals-{stamp}.csv"'
        # BOM تا اکسل فارسی را درست باز کند؛ بدون آن ستون‌ها به‌هم می‌ریزند.
        response.write('﻿')
        writer = csv.writer(response)
        writer.writerow(EXPORT_COLUMNS)
        for row in rows:
            writer.writerow(row)
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'پرداخت برداشت'
    ws.append(EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color='D4AF37', end_color='D4AF37', fill_type='solid')
    for row in rows:
        ws.append(row)
    # شبا باید متن بماند، وگرنه اکسل آن را عدد می‌کند و صفرهای ابتدایی می‌پرند.
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = '@'
        ws.cell(row=r, column=4).number_format = '#,##0'
    for col, width in zip('ABCDEFGHI', (14, 26, 30, 16, 26, 15, 14, 18, 22)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="withdrawals-{stamp}.xlsx"'
    wb.save(response)
    return response


@staff_member_required
def admin_withdrawal_import(request):
    """همان فایل را با ستونِ «شماره پیگیری واریز» پرشده برمی‌گرداند.

    فقط سطرهایی که شماره پیگیری دارند پردازش می‌شوند. مبلغِ داخل فایل با
    رکورد مقایسه می‌شود: اگر یکی نبود، آن سطر رد می‌شود -- یعنی خزانه‌دار
    مبلغ دیگری پرداخت کرده و ثبت خودکارش وضعیت را غلط نشان می‌داد.
    """
    if request.method != 'POST' or not request.FILES.get('file'):
        messages.error(request, 'فایلی انتخاب نشده است.')
        return redirect('wallet:admin_withdrawal_list')

    upload = request.FILES['file']
    try:
        rows = _read_uploaded_rows(upload)
    except Exception as exc:                            # noqa: BLE001
        messages.error(request, f'فایل خوانده نشد: {exc}')
        return redirect('wallet:admin_withdrawal_list')

    paid = skipped = 0
    problems = []
    for lineno, (req_id, amount, bank_ref) in enumerate(rows, start=2):
        if not bank_ref:
            continue                                    # هنوز پرداخت نشده
        wr = WithdrawalRequest.objects.filter(pk=req_id).first()
        if wr is None:
            problems.append(f'سطر {lineno}: درخواست {req_id} پیدا نشد')
            skipped += 1
            continue
        if amount is not None and amount != wr.amount:
            problems.append(
                f'سطر {lineno}: مبلغ فایل ({amount:,}) با درخواست #{req_id} '
                f'({wr.amount:,}) یکی نیست')
            skipped += 1
            continue
        ok, err = wr.mark_paid(request.user, bank_reference=bank_ref,
                               note='ثبت گروهی از فایل خزانه‌داری')
        if ok:
            paid += 1
        else:
            problems.append(f'سطر {lineno}: درخواست #{req_id} — {err}')
            skipped += 1

    if paid:
        messages.success(request, f'{paid} واریز از روی فایل ثبت شد.')
    if problems:
        messages.warning(request, f'{skipped} سطر ثبت نشد: ' + ' | '.join(problems[:10]))
    if not paid and not problems:
        messages.info(request, 'هیچ سطری با شماره پیگیری پرشده در فایل نبود.')
    return redirect('wallet:admin_withdrawal_list')


def _read_uploaded_rows(upload):
    """(شماره درخواست، مبلغ، شماره پیگیری) را از xlsx یا csv بیرون می‌کشد."""
    digits = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')

    def to_int(v):
        if v in (None, ''):
            return None
        s = str(v).translate(digits).replace(',', '').replace('٬', '').strip()
        return int(float(s)) if s else None

    name = (upload.name or '').lower()
    raw = []
    if name.endswith('.csv'):
        import csv
        import io
        text = upload.read().decode('utf-8-sig', errors='replace')
        raw = list(csv.reader(io.StringIO(text)))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(upload, data_only=True)
        raw = [[c for c in row] for row in wb.active.iter_rows(values_only=True)]

    out = []
    for row in raw[1:]:                                 # سطر اول عنوان ستون‌هاست
        if not row or all(c in (None, '') for c in row):
            continue
        req_id = to_int(row[0] if len(row) > 0 else None)
        amount = to_int(row[3] if len(row) > 3 else None)
        bank_ref = str(row[8]).strip() if len(row) > 8 and row[8] not in (None, '') else ''
        if req_id is None:
            continue
        out.append((req_id, amount, bank_ref))
    return out